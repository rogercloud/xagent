import errno
import json
import time
from pathlib import Path
from unittest.mock import Mock

import pytest

from xagent.web.tools.mcp import google_slides


@pytest.fixture(autouse=True)
def _credentials(monkeypatch):
    # Every test below replaces get_slides_service() wholesale via
    # _mock_slides_service, so nothing here exercises real credential
    # building — this just guards against a stray, unmocked call blowing up
    # with a confusing "missing env var" error instead of the actual assertion.
    monkeypatch.setenv("GOOGLE_ACCESS_TOKEN", "access-token")
    monkeypatch.setattr(google_slides, "_CREATED_DEFAULT_SLIDES", {})
    monkeypatch.setattr(google_slides, "_PRESERVED_BLANK_SLIDES", {})


def _mock_slides_service(monkeypatch, presentations_mock):
    service = Mock()
    service.presentations.return_value = presentations_mock
    if not isinstance(presentations_mock.get.return_value.execute.return_value, dict):
        presentations_mock.get.return_value.execute.return_value = {"slides": []}
    monkeypatch.setattr(google_slides, "get_slides_service", lambda: service)
    return service


def _batch_update_requests(presentations_mock):
    return presentations_mock.batchUpdate.call_args.kwargs["body"]["requests"]


def _placeholder_object_id(create_slide, placeholder_type):
    return next(
        m["objectId"]
        for m in create_slide["placeholderIdMappings"]
        if m["layoutPlaceholder"]["type"] == placeholder_type
    )


def _placeholder_element(object_id, placeholder_type, text=""):
    text_elements = [{"textRun": {"content": text}}] if text else []
    return {
        "objectId": object_id,
        "shape": {
            "placeholder": {"type": placeholder_type},
            "text": {"textElements": text_elements},
        },
    }


def _mock_presentation_get(presentations_mock, slide_id, elements):
    presentations_mock.get.return_value.execute.return_value = {
        "slides": [{"objectId": slide_id, "pageElements": elements}]
    }


def _mock_pptx_text(monkeypatch, *slide_texts, has_content=None):
    content_flags = (
        list(has_content)
        if has_content is not None
        else [bool(text.strip()) for text in slide_texts]
    )
    assert len(content_flags) == len(slide_texts)
    monkeypatch.setattr(
        google_slides,
        "_extract_pptx_slide_content",
        lambda _file_path: list(zip(slide_texts, content_flags, strict=True)),
    )


def test_add_slide_default_layout_creates_title_and_body_with_bullets(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Q3 Pipeline Highlights", body="Line one\nLine two"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)

    create_slide = requests[0]["createSlide"]
    assert create_slide["slideLayoutReference"] == {
        "predefinedLayout": "TITLE_AND_BODY"
    }
    mappings = {
        m["layoutPlaceholder"]["type"]: m["objectId"]
        for m in create_slide["placeholderIdMappings"]
    }
    assert set(mappings) == {"TITLE", "BODY"}

    title_req = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == mappings["TITLE"]
    )
    assert title_req["text"] == "Q3 Pipeline Highlights"
    body_req = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == mappings["BODY"]
    )
    assert body_req["text"] == "Line one\nLine two"

    bullets_req = next(
        r["createParagraphBullets"] for r in requests if "createParagraphBullets" in r
    )
    assert bullets_req["objectId"] == mappings["BODY"]
    assert bullets_req["textRange"] == {"type": "ALL"}
    assert bullets_req["bulletPreset"] == "BULLET_DISC_CIRCLE_SQUARE"

    # createParagraphBullets must come after the insertText that populates
    # the body — applying it to an empty range fails against the real API.
    body_insert_index = next(
        i
        for i, r in enumerate(requests)
        if "insertText" in r and r["insertText"]["objectId"] == mappings["BODY"]
    )
    bullets_index = next(
        i for i, r in enumerate(requests) if "createParagraphBullets" in r
    )
    assert bullets_index > body_insert_index


def test_add_slide_strips_literal_bullet_markers_before_inserting(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    body = "• First point\n- Second point\n* Third point\nFourth point (no marker)"
    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body=body)
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == (
        "First point\nSecond point\nThird point\nFourth point (no marker)"
    )


def test_add_slide_strips_marker_glued_to_text_without_trailing_space(monkeypatch):
    """A marker with no space after it (e.g. "•First", "-Nospace") must
    still be recognized — otherwise the literal marker survives next to
    Slides' own bullet glyph, a visible double bullet."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", body="•First point\n-Nospace"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "First point\nNospace"


def test_add_slide_preserves_nested_bullet_indentation(monkeypatch):
    """Slides infers bullet nesting level from leading whitespace/tabs in
    the inserted text; stripping the marker must not also strip the
    indentation in front of it, or multi-level bullets flatten to one
    level."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    body = "- Top\n  - Nested\n    - Deeper"
    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body=body)
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    # Slides infers nesting level from leading *tabs*, not spaces — plain
    # spaces would just be inserted as literal, non-nesting text — so
    # 2-space indentation must be converted to real tab characters, one
    # tab per level, for the nesting to actually render in Slides.
    assert body_text == "Top\n\tNested\n\t\tDeeper"


def test_add_slide_preserves_literal_tab_indentation(monkeypatch):
    """A caller who already typed real tab characters (rather than
    spaces) for nesting must have them preserved 1:1 — this is the input
    shape Slides' createParagraphBullets actually keys off."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    body = "- Top\n\t- Nested\n\t\t- Deeper"
    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body=body)
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "Top\n\tNested\n\t\tDeeper"


def test_add_slide_mixed_tab_and_space_indentation_uses_tab_count_only(monkeypatch):
    """Documented limitation: mixing tabs and spaces in one line's leading
    whitespace is not combined — if any tab is present, the tab count
    wins and spaces in that same run are ignored."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    body = "- Top\n\t  - Mixed"
    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body=body)
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "Top\n\tMixed"


def test_add_slide_does_not_corrupt_non_bullet_content_starting_with_marker_chars(
    monkeypatch,
):
    """ "-"/"*" have common non-bullet meanings (a negative number's sign,
    a currency figure, a decimal, an em-/en-dash, markdown emphasis) — a
    leading marker glued to a digit, symbol, or another marker character,
    or a bare "*" with no trailing space, must be left alone rather than
    silently mangled. This is a regression guard for a real bug: an
    earlier fix for "•First" (marker glued to a word) also silently
    stripped the sign off negative figures like "-$5m loss"."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    body = (
        "-5% growth\n-$5m loss\n-.5% decline\n-- Author Name\n-– Author\n"
        "**Note\n*emphasis* not a bullet"
    )
    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body=body)
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == body


def test_add_slide_strips_bare_marker_with_nothing_after_it(monkeypatch):
    """A line that's only a bullet marker with nothing after it at all
    (e.g. a stray "•" or "-" on its own line) previously wasn't recognized
    by the stripping regex (which required a character or whitespace
    after the marker), so it survived as literal text right next to
    Slides' own bullet glyph — a visible double bullet."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", body="Point one\n•\nPoint two"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "Point one\nPoint two"


def test_add_slide_fully_unwraps_a_line_with_more_than_one_leading_marker(monkeypatch):
    """ "• - nested" must not leave a residual "-" as literal text inside a
    paragraph Slides is about to bullet — strip leading markers repeatedly
    until none remain, not just the outermost one."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide("pres1", title="T", body="• - nested")

    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "nested"


def test_strip_bullet_prefixes_bounds_work_on_a_degenerate_glued_marker_run():
    """Regression guard for a real perf/DoS finding: _strip_line used to
    re-scan the whole remaining line on every stripped marker with no
    cap, making a line of N glued "•" characters (body has no length
    limit; plausible from a runaway/malformed LLM completion) O(n^2) —
    measured at over a second for N=200_000. _MAX_MARKER_STRIPS_PER_LINE
    must keep this bounded regardless of N."""
    line = "•" * 200_000 + "Hello"

    start = time.perf_counter()
    result = google_slides._strip_bullet_prefixes(line)
    elapsed = time.perf_counter() - start

    assert elapsed < 1.0
    assert result.endswith("Hello")


def test_add_slide_drops_blank_lines_and_trailing_newline_from_bulleted_body(
    monkeypatch,
):
    """A blank line inside body, or a trailing newline (very common in
    generated text), would otherwise become an empty paragraph that still
    gets createParagraphBullets applied to it — a visible floating bullet
    with no text next to it."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide(
        "pres1", title="T", body="Point one\n\nPoint two\n"
    )

    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert body_text == "Point one\nPoint two"


def test_add_slide_rejects_marker_only_body_for_content_layout(monkeypatch):
    """A body that's only a bullet marker ("•   ") strips to an empty
    string before insertion — the body-required guard must check the
    post-stripping text, not the raw string, or this recreates the
    content-less-slide bug via a different input shape."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", body="•   ", layout="TITLE_AND_BODY"
        )
    )

    assert result["status"] == "error"
    assert "expects body content but none was provided" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_title_layout_uses_subtitle_and_skips_bullets(monkeypatch):
    """The TITLE (cover) layout maps body to a SUBTITLE placeholder, which
    should not get a createParagraphBullets request — a subtitle line isn't a
    bulleted list."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Q3 2026 Sales & CRM Review",
            body="Pipeline Updates, Key Wins, and Q4 Strategy",
            layout="TITLE",
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)

    create_slide = requests[0]["createSlide"]
    assert create_slide["slideLayoutReference"] == {"predefinedLayout": "TITLE"}
    mappings = {
        m["layoutPlaceholder"]["type"]: m["objectId"]
        for m in create_slide["placeholderIdMappings"]
    }
    assert set(mappings) == {"CENTERED_TITLE", "SUBTITLE"}
    assert not any("createParagraphBullets" in r for r in requests)


@pytest.mark.parametrize("layout", ["TITLE_ONLY", "SECTION_HEADER"])
def test_add_slide_title_only_layouts_need_no_body(monkeypatch, layout):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Section Break", layout=layout
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    create_slide = requests[0]["createSlide"]
    assert create_slide["slideLayoutReference"] == {"predefinedLayout": layout}
    mappings = {
        m["layoutPlaceholder"]["type"] for m in create_slide["placeholderIdMappings"]
    }
    assert mappings == {"TITLE"}


@pytest.mark.parametrize("layout", ["TITLE_ONLY", "SECTION_HEADER"])
def test_add_slide_rejects_empty_title_for_title_only_layout(monkeypatch, layout):
    """Regression guard, symmetric with the body-required check: a layout
    whose only content slot is the title must not silently create a fully
    empty slide when title is also omitted."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_add_slide("pres1", layout=layout))

    assert result["status"] == "error"
    assert "completely empty slide" in result["message"]
    presentations.batchUpdate.assert_not_called()


@pytest.mark.parametrize("layout", ["TITLE_ONLY", "SECTION_HEADER"])
def test_add_slide_rejects_title_that_is_only_a_newline(monkeypatch, layout):
    """Regression guard: before the newline-collapse + .strip()-based
    guards were added together, title="\\n" was truthy under the old bare
    truthiness check and slipped past the empty-slide guard, producing a
    createSlide with no insertText for the title at all (insertion always
    used title.strip()) — a silently blank slide. Collapsing "\\n" to " "
    and then stripping it must route this through the empty-slide guard
    with its real message, not a misleading one from elsewhere."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="\n", layout=layout)
    )

    assert result["status"] == "error"
    assert "completely empty slide" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_rejects_completely_empty_title_layout(monkeypatch):
    """The TITLE (cover) layout accepts an empty body (subtitle is
    optional), but title and body can't both be empty — that's a
    completely blank cover slide."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_add_slide("pres1", layout="TITLE"))

    assert result["status"] == "error"
    assert "completely empty slide" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_skips_whitespace_only_title_insertion(monkeypatch):
    """A whitespace-only title on a content layout (where body carries the
    real content) must not be inserted verbatim — leave the placeholder
    empty instead of filling it with whitespace."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide(
        "pres1", title="   ", body="Real content", layout="TITLE_AND_BODY"
    )

    requests = _batch_update_requests(presentations)
    title_object_id = _placeholder_object_id(requests[0]["createSlide"], "TITLE")
    assert not any(
        "insertText" in r and r["insertText"]["objectId"] == title_object_id
        for r in requests
    )


def test_add_slide_rejects_unknown_layout(monkeypatch):
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", layout="TWO_COLUMNS")
    )

    assert result["status"] == "error"
    assert "TWO_COLUMNS" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_rejects_non_string_layout_from_direct_call(monkeypatch):
    """layout's Literal typing is only enforced by FastMCP's validation
    layer, which a direct Python call bypasses entirely — a non-string
    value (e.g. an int) must still surface as a clean error rather than
    an unhandled AttributeError from calling .strip() on it."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", layout=123)
    )

    assert result["status"] == "error"
    assert "'layout' must be a string" in result["message"]
    presentations.batchUpdate.assert_not_called()


@pytest.mark.parametrize("layout", ["TITLE_ONLY", "SECTION_HEADER"])
def test_add_slide_rejects_body_on_layout_without_body_placeholder(monkeypatch, layout):
    """Regression guard: these layouts have no body placeholder, so silently
    accepting `body` would drop it exactly like the reported bug — reject the
    call instead so the caller finds out immediately."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", body="This would be lost", layout=layout
        )
    )

    assert result["status"] == "error"
    assert "has no body placeholder" in result["message"]
    presentations.batchUpdate.assert_not_called()


@pytest.mark.parametrize("layout", ["TITLE_ONLY", "SECTION_HEADER"])
def test_add_slide_allows_whitespace_only_body_on_layout_without_body_placeholder(
    monkeypatch, layout
):
    """Regression guard: the "no body placeholder" rejection must key off
    stripped content, matching insertion (which already treats a
    whitespace-only body as absent) — otherwise body="\\n" is hard-rejected
    even though nothing would actually have been dropped."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", body="\n", layout=layout
        )
    )

    assert result["status"] == "success"


@pytest.mark.parametrize(
    "kwargs", [{"title": "   "}, {"title": "T"}, {"body": "B"}, {}]
)
def test_add_slide_rejects_blank_layout(monkeypatch, kwargs):
    """Blank/custom slides must use batch_update rather than this primitive."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", layout="BLANK", **kwargs)
    )

    assert result["status"] == "error"
    assert "Unknown layout" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_rejects_missing_body_for_content_layout(monkeypatch):
    """Regression guard for the reported bug: a content layout (a real BODY
    placeholder) must not silently create a slide with no detail text."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Q3 Pipeline Highlights", layout="TITLE_AND_BODY"
        )
    )

    assert result["status"] == "error"
    assert "expects body content but none was provided" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_rejects_whitespace_only_body_for_content_layout(monkeypatch):
    """Regression guard: a whitespace-only body ("   ", "\\n\\n") must not
    slip past the "body required" check just because it's non-empty — that
    would recreate the exact content-less-slide bug the check exists for."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Q3 Pipeline Highlights",
            body="   \n\n  ",
            layout="TITLE_AND_BODY",
        )
    )

    assert result["status"] == "error"
    assert "expects body content but none was provided" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_add_slide_title_layout_preserves_literal_dash_in_subtitle(monkeypatch):
    """Regression guard: bullet-marker stripping must only apply to text
    that's actually being turned into a bulleted list (a real BODY
    placeholder) — a SUBTITLE is never bulleted, so a literal leading "-"
    the caller intended as part of the subtitle text must survive."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Guide", body="- The Complete Guide", layout="TITLE"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    subtitle_object_id = _placeholder_object_id(requests[0]["createSlide"], "SUBTITLE")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == subtitle_object_id
    )
    assert body_text == "- The Complete Guide"


def test_add_slide_drops_blank_lines_on_non_bulleted_body(monkeypatch):
    """Blank-line filtering must not be limited to bulleted content — a
    blank line or trailing newline in a non-bulleted body (e.g. TITLE's
    subtitle) would otherwise leave stray blank paragraphs in the
    placeholder even though there's no floating-bullet symptom to notice."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Guide", body="\n\nSubtitle\n", layout="TITLE"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    subtitle_object_id = _placeholder_object_id(requests[0]["createSlide"], "SUBTITLE")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == subtitle_object_id
    )
    assert body_text == "Subtitle"


def test_add_slide_skips_whitespace_only_body_insertion_on_title_layout(monkeypatch):
    """Regression guard, symmetric with the whitespace-only-title fix: a
    whitespace-only body on a layout where body is optional (e.g. TITLE's
    subtitle) must not be inserted verbatim — leave the placeholder
    untouched instead of writing invisible whitespace into it."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Cover", body="   ", layout="TITLE"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    subtitle_object_id = _placeholder_object_id(requests[0]["createSlide"], "SUBTITLE")
    assert not any(
        "insertText" in r and r["insertText"]["objectId"] == subtitle_object_id
        for r in requests
    )


def test_add_slide_title_layout_allows_missing_body(monkeypatch):
    """A cover slide's subtitle is optional, unlike a real BODY placeholder."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Cover", layout="TITLE")
    )

    assert result["status"] == "success"


def test_add_slide_resolves_full_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    url = "https://docs.google.com/presentation/d/abc123/edit#slide=id.p"
    result = json.loads(
        google_slides.google_slides_add_slide(url, title="T", body="detail line")
    )

    assert result["status"] == "success"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "abc123"


def test_add_slide_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body="detail")
    )

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_add_slide_returns_error_payload_without_ascii_escaping(monkeypatch):
    """_error() must use ensure_ascii=False like every success payload in
    this file, or non-ASCII error text (e.g. from a caller's own input
    echoed back, or a non-ASCII exception message) gets mangled into
    \\uXXXX escapes instead of staying human-readable."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError(
        "读取失败 — 😀"
    )
    _mock_slides_service(monkeypatch, presentations)

    raw = google_slides.google_slides_add_slide("pres1", title="T", body="detail")

    assert "\\u" not in raw
    assert "读取失败" in raw
    assert json.loads(raw)["status"] == "error"


@pytest.mark.parametrize(
    ("layout", "kwargs", "expected_layout"),
    [
        ("title_and_body", {"body": "detail"}, "TITLE_AND_BODY"),
        (" TITLE ", {}, "TITLE"),
        ("Title_Only", {}, "TITLE_ONLY"),
    ],
)
def test_add_slide_normalizes_layout_case_and_whitespace(
    monkeypatch, layout, kwargs, expected_layout
):
    """An LLM caller may not reproduce the exact enum casing/spacing —
    normalize before rejecting as unknown."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="T", layout=layout, **kwargs
        )
    )

    assert result["status"] == "success"
    assert result["layout"] == expected_layout


def test_add_slide_echoes_the_effective_layout_on_success(monkeypatch):
    """The caller passed no explicit layout, relying on the default — the
    response should confirm which layout was actually applied."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body="detail")
    )

    assert result["layout"] == "TITLE_AND_BODY"


def test_add_slide_strips_padding_whitespace_from_title_before_inserting(monkeypatch):
    """Unlike body (whose leading whitespace is meaningful for bullet
    nesting), a title has no reason to keep incidental leading/trailing
    padding — insert the trimmed text, not the raw string."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide("pres1", title="  Q3 Review  ", body="detail")

    requests = _batch_update_requests(presentations)
    title_object_id = _placeholder_object_id(requests[0]["createSlide"], "TITLE")
    title_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == title_object_id
    )
    assert title_text == "Q3 Review"


def test_add_slide_collapses_embedded_newline_in_title(monkeypatch):
    """A title is expected to be a single line — an embedded newline
    (e.g. from a caller accidentally pasting multi-line content into
    title) must be collapsed to a space rather than silently producing a
    multi-paragraph title placeholder."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide("pres1", title="Q3\nReview", body="detail")

    requests = _batch_update_requests(presentations)
    title_object_id = _placeholder_object_id(requests[0]["createSlide"], "TITLE")
    title_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == title_object_id
    )
    assert title_text == "Q3 Review"


def test_add_slide_normalizes_crlf_and_cr_newlines(monkeypatch):
    """Text pasted from a Windows editor may use \\r\\n or lone \\r line
    endings; leaving stray \\r characters embedded in the inserted text is
    a latent artifact even though it doesn't break bullet-marker matching
    (which only anchors on line starts)."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    google_slides.google_slides_add_slide(
        "pres1", title="T", body="Line one\r\nLine two\rLine three"
    )

    requests = _batch_update_requests(presentations)
    body_object_id = _placeholder_object_id(requests[0]["createSlide"], "BODY")
    body_text = next(
        r["insertText"]["text"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == body_object_id
    )
    assert "\r" not in body_text
    assert body_text == "Line one\nLine two\nLine three"


async def test_add_slide_rejects_unknown_layout_via_mcp_layer(monkeypatch):
    """layout is typed as a Literal enum so FastMCP validates it before the
    function body ever runs — exercise the real call path, which direct
    function calls (used by every other test in this file) bypass."""
    from mcp.server.fastmcp.exceptions import ToolError

    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    with pytest.raises(ToolError, match="validation error"):
        await google_slides.mcp.call_tool(
            "google_slides_add_slide",
            {"presentation_id": "pres1", "title": "T", "layout": "TWO_COLUMNS"},
        )
    presentations.batchUpdate.assert_not_called()


async def test_add_slide_normalizes_layout_via_mcp_layer(monkeypatch):
    """Regression guard: `layout`'s Literal type is validated by FastMCP's
    Pydantic layer *before* the function body runs — a naive Literal
    annotation would reject a differently-cased/spaced value there,
    making the function's own layout.strip().upper() dead code for every
    real (non-test, non-direct-call) caller. `_normalize_layout` must run
    as a BeforeValidator so normalization happens ahead of the Literal
    check, not after it."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    content, _ = await google_slides.mcp.call_tool(
        "google_slides_add_slide",
        {
            "presentation_id": "pres1",
            "title": "T",
            "body": "detail",
            "layout": " title_and_body ",
        },
    )

    result = json.loads(content[0].text)
    assert result["status"] == "success"
    assert result["layout"] == "TITLE_AND_BODY"


async def test_update_slide_via_mcp_layer(monkeypatch):
    """Smoke test through the real MCP dispatch path (JSON args dict,
    positional/keyword binding via call_tool), not just a direct Python
    call — the only path add_slide's Literal-normalization bug (fixed
    above) was actually reachable through."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE")]
    )

    content, _ = await google_slides.mcp.call_tool(
        "google_slides_update_slide",
        {"presentation_id": "pres1", "slide_id": "slide1", "title": "New title"},
    )

    result = json.loads(content[0].text)
    assert result["status"] == "success"


async def test_delete_slide_via_mcp_layer(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(presentations, "slide1", [])

    content, _ = await google_slides.mcp.call_tool(
        "google_slides_delete_slide",
        {"presentation_id": "pres1", "slide_id": "slide1"},
    )

    result = json.loads(content[0].text)
    assert result["status"] == "success"


def test_update_slide_replaces_title_and_body_and_reapplies_bullets(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [
            _placeholder_element("title_obj", "TITLE", text="Old title"),
            _placeholder_element("body_obj", "BODY", text="Old body"),
        ],
    )

    result = json.loads(
        google_slides.google_slides_update_slide(
            "pres1", "slide1", title="New title", body="• Fixed detail"
        )
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)

    delete_ids = [r["deleteText"]["objectId"] for r in requests if "deleteText" in r]
    assert set(delete_ids) == {"title_obj", "body_obj"}

    title_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "title_obj"
    )
    assert title_insert["text"] == "New title"

    body_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "body_obj"
    )
    assert body_insert["text"] == "Fixed detail"

    bullets_req = next(
        r["createParagraphBullets"] for r in requests if "createParagraphBullets" in r
    )
    assert bullets_req["objectId"] == "body_obj"

    # Order matters: batchUpdate applies requests in list order, so a
    # deleteText must precede the insertText for the same objectId (an
    # insert-then-delete would wipe out the new text instead of the old),
    # and createParagraphBullets must come after the body insertText it
    # formats.
    def _index_of(predicate):
        return next(i for i, r in enumerate(requests) if predicate(r))

    title_delete_index = _index_of(
        lambda r: r.get("deleteText", {}).get("objectId") == "title_obj"
    )
    title_insert_index = _index_of(
        lambda r: r.get("insertText", {}).get("objectId") == "title_obj"
    )
    assert title_delete_index < title_insert_index

    body_delete_index = _index_of(
        lambda r: r.get("deleteText", {}).get("objectId") == "body_obj"
    )
    body_insert_index = _index_of(
        lambda r: r.get("insertText", {}).get("objectId") == "body_obj"
    )
    bullets_index = _index_of(lambda r: "createParagraphBullets" in r)
    assert body_delete_index < body_insert_index < bullets_index


def test_update_slide_skips_delete_text_when_placeholder_already_empty(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("title_obj", "TITLE", text="")],
    )

    google_slides.google_slides_update_slide("pres1", "slide1", title="First title")

    requests = _batch_update_requests(presentations)
    assert not any("deleteText" in r for r in requests)
    assert requests[0]["insertText"]["text"] == "First title"


def test_update_slide_skips_delete_text_when_placeholder_has_only_a_newline(
    monkeypatch,
):
    """Regression guard: the Slides API commonly represents a "cleared"
    placeholder as a lone trailing "\\n" (the paragraph terminator), not a
    truly empty string — the existing-text check must treat that the same
    as empty, or it triggers a needless deleteText."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("title_obj", "TITLE", text="\n")],
    )

    google_slides.google_slides_update_slide("pres1", "slide1", title="First title")

    requests = _batch_update_requests(presentations)
    assert not any("deleteText" in r for r in requests)
    assert requests[0]["insertText"]["text"] == "First title"


def test_update_slide_still_clears_whitespace_beyond_the_bare_terminator(monkeypatch):
    """Regression guard: only the exact "" or "\\n" (the implicit
    terminator) should skip deleteText — any other whitespace-only text
    (e.g. a stray " \\n" left by a slide edited by something other than
    this tool) must still be cleared. insertText has no insertionIndex
    here, so it defaults to prepending rather than replacing; skipping
    the delete for arbitrary whitespace would leave that stale text
    merged into what's supposed to be a clean replacement."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("title_obj", "TITLE", text=" \n")],
    )

    google_slides.google_slides_update_slide("pres1", "slide1", title="First title")

    requests = _batch_update_requests(presentations)
    delete_req = next(r["deleteText"] for r in requests if "deleteText" in r)
    assert delete_req["objectId"] == "title_obj"


def test_update_slide_only_touches_the_field_provided(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [
            _placeholder_element("title_obj", "TITLE", text="Old title"),
            _placeholder_element("body_obj", "BODY", text="Old body"),
        ],
    )

    google_slides.google_slides_update_slide("pres1", "slide1", title="New title")

    requests = _batch_update_requests(presentations)
    assert not any(
        r.get("deleteText", {}).get("objectId") == "body_obj"
        or r.get("insertText", {}).get("objectId") == "body_obj"
        for r in requests
    )


def test_update_slide_subtitle_body_is_not_bulleted_or_stripped(monkeypatch):
    """A TITLE-layout slide's body lands in a SUBTITLE, not a BODY —
    update_slide must follow the same non-bulleted, non-stripped rule as
    add_slide for that placeholder type."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [
            _placeholder_element("title_obj", "CENTERED_TITLE", text="Old"),
            _placeholder_element("subtitle_obj", "SUBTITLE", text="Old subtitle"),
        ],
    )

    google_slides.google_slides_update_slide(
        "pres1", "slide1", body="- Literal dash subtitle"
    )

    requests = _batch_update_requests(presentations)
    body_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "subtitle_obj"
    )
    assert body_insert["text"] == "- Literal dash subtitle"
    assert not any("createParagraphBullets" in r for r in requests)


def test_update_slide_requires_at_least_one_field(monkeypatch):
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_update_slide("pres1", "slide1"))

    assert result["status"] == "error"
    presentations.get.assert_not_called()


def test_update_slide_rejects_whitespace_only_title(monkeypatch):
    """Regression guard, symmetric with the body-side check above:
    whitespace-only title text ("   ") must not slip past the guard just
    because it's non-empty."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="   ")
    )

    assert result["status"] == "error"
    presentations.get.assert_not_called()


def test_update_slide_rejects_whitespace_only_body(monkeypatch):
    """Regression guard: whitespace-only text ("   ", "\\n\\n") must not
    slip past the guard just because it's non-empty — that would silently
    blank real slide content, the same bug class fixed for add_slide."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", body="   \n\n  ")
    )

    assert result["status"] == "error"
    presentations.get.assert_not_called()


def test_update_slide_rejects_marker_only_body(monkeypatch):
    """Regression guard, same class as add_slide's: a body that's only a
    bullet marker ("•   ") passes the raw whitespace-only check (it's
    non-blank) but strips down to nothing once bullet-marker/blank-line
    processing runs — must still be rejected, not silently applied as an
    empty update."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("body_obj", "BODY", text="Old body")],
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", body="•   ")
    )

    assert result["status"] == "error"
    presentations.batchUpdate.assert_not_called()


def test_update_slide_rejects_unknown_slide_id(monkeypatch):
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(presentations, "other_slide", [])

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="T")
    )

    assert result["status"] == "error"
    assert "slide1" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_update_slide_rejects_title_when_slide_has_no_title_placeholder(monkeypatch):
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("body_obj", "BODY", text="x")]
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="T")
    )

    assert result["status"] == "error"
    assert "title" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_update_slide_rejects_body_when_slide_has_no_body_placeholder(monkeypatch):
    """Regression guard, symmetric with the title-side test above."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE", text="x")]
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", body="Detail text")
    )

    assert result["status"] == "error"
    assert "body" in result["message"]
    presentations.batchUpdate.assert_not_called()


def test_update_slide_collapses_embedded_newline_in_title(monkeypatch):
    """Regression guard, symmetric with add_slide's fix: a title is a
    single-line placeholder — an embedded newline must collapse to a
    space instead of silently producing a multi-paragraph title."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE")]
    )

    google_slides.google_slides_update_slide(
        "pres1", "slide1", title="Line one\nLine two"
    )

    requests = _batch_update_requests(presentations)
    title_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "title_obj"
    )
    assert title_insert["text"] == "Line one Line two"


def test_update_slide_normalizes_crlf_and_cr_newlines(monkeypatch):
    """Regression guard, symmetric with add_slide's fix: text pasted from
    a Windows editor may use \\r\\n or lone \\r line endings — these must
    not survive as stray \\r characters embedded in the inserted text."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("body_obj", "BODY")],
    )

    google_slides.google_slides_update_slide(
        "pres1", "slide1", body="Line one\r\nLine two\rLine three"
    )

    requests = _batch_update_requests(presentations)
    body_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "body_obj"
    )
    assert "\r" not in body_insert["text"]
    assert body_insert["text"] == "Line one\nLine two\nLine three"


def test_update_slide_rejects_whole_call_when_body_becomes_empty_even_with_title(
    monkeypatch,
):
    """Regression guard: a marker-only body must reject the entire call
    (no batchUpdate at all), not silently apply just the title update and
    drop the body — the two are meant to be one atomic edit."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [
            _placeholder_element("title_obj", "TITLE", text="Old title"),
            _placeholder_element("body_obj", "BODY", text="Old body"),
        ],
    )

    result = json.loads(
        google_slides.google_slides_update_slide(
            "pres1", "slide1", title="New title", body="•   "
        )
    )

    assert result["status"] == "error"
    presentations.batchUpdate.assert_not_called()


def test_update_slide_writes_title_into_centered_title_placeholder(monkeypatch):
    """A TITLE-layout slide's title lands in a CENTERED_TITLE placeholder,
    not a plain TITLE — update_slide must write to it just the same."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("title_obj", "CENTERED_TITLE", text="Old")],
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="New title")
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    title_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "title_obj"
    )
    assert title_insert["text"] == "New title"


def test_update_slide_writes_body_into_object_placeholder(monkeypatch):
    """OBJECT is a real Slides placeholder type common on slides from an
    imported/non-standard-theme presentation (e.g. one converted from
    PowerPoint) — update_slide must recognize it as a body-role
    placeholder like BODY/SUBTITLE, not reject it as unrecognized."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("body_obj", "OBJECT", text="Old detail")],
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", body="New detail")
    )

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    body_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "body_obj"
    )
    assert body_insert["text"] == "New detail"
    # OBJECT is a generic content placeholder, not the same as a real
    # bulleted BODY list — bullets are only applied for the "BODY" type.
    assert not any("createParagraphBullets" in r for r in requests)


def test_update_slide_only_targets_the_first_placeholder_of_a_duplicated_role(
    monkeypatch,
):
    """Documented, accepted limitation: a slide with two BODY placeholders
    (not reachable via this file's own add_slide, but possible for a slide
    from another source) silently targets only the first one found. Lock
    in the current behavior so a future change to iteration order isn't
    silent."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [
            _placeholder_element("body_obj_1", "BODY", text="First"),
            _placeholder_element("body_obj_2", "BODY", text="Second"),
        ],
    )

    google_slides.google_slides_update_slide("pres1", "slide1", body="New detail")

    requests = _batch_update_requests(presentations)
    assert not any(
        r.get("insertText", {}).get("objectId") == "body_obj_2"
        or r.get("deleteText", {}).get("objectId") == "body_obj_2"
        for r in requests
    )
    body_insert = next(
        r["insertText"]
        for r in requests
        if "insertText" in r and r["insertText"]["objectId"] == "body_obj_1"
    )
    assert body_insert["text"] == "New detail"


def test_update_slide_resolves_full_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE")]
    )

    url = "https://docs.google.com/presentation/d/abc123/edit#slide=id.p"
    result = json.loads(
        google_slides.google_slides_update_slide(url, "slide1", title="T")
    )

    assert result["status"] == "success"
    assert presentations.get.call_args.kwargs["presentationId"] == "abc123"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "abc123"


def test_update_slide_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="T")
    )

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_update_slide_returns_error_payload_on_batch_update_failure(monkeypatch):
    """Symmetric with the get() failure case above and with delete_slide's
    own batchUpdate-failure test — the earlier presentations().get() call
    can succeed while the actual edit still fails."""
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE")]
    )

    result = json.loads(
        google_slides.google_slides_update_slide("pres1", "slide1", title="T")
    )

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_delete_slide_sends_delete_object_request(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(presentations, "slide1", [])

    result = json.loads(google_slides.google_slides_delete_slide("pres1", "slide1"))

    assert result["status"] == "success"
    requests = _batch_update_requests(presentations)
    assert requests == [{"deleteObject": {"objectId": "slide1"}}]


def test_delete_slide_rejects_id_that_is_not_a_slide(monkeypatch):
    """Regression guard: a placeholder shape id (e.g. one this file itself
    mints as f"{slide_id}_title") must not be silently accepted — Slides'
    deleteObject would delete just that shape while reporting success as if
    the whole slide had been removed. The mocked presentation genuinely
    contains a shape with this id — nested inside slide1's pageElements,
    not as a top-level slide — confirming _find_slide's rejection comes
    from checking only top-level slide ids (by design), not from the id
    being absent from the API response altogether."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("slide1_title", "TITLE", text="Real title")],
    )

    result = json.loads(
        google_slides.google_slides_delete_slide("pres1", "slide1_title")
    )

    assert result["status"] == "error"
    presentations.batchUpdate.assert_not_called()


def test_delete_slide_resolves_full_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(presentations, "slide1", [])

    url = "https://docs.google.com/presentation/d/abc123/edit#slide=id.p"
    result = json.loads(google_slides.google_slides_delete_slide(url, "slide1"))

    assert result["status"] == "success"
    assert presentations.get.call_args.kwargs["presentationId"] == "abc123"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "abc123"


def test_delete_slide_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(presentations, "slide1", [])

    result = json.loads(google_slides.google_slides_delete_slide("pres1", "slide1"))

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_get_presentation_returns_slide_summaries(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "My Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {
                        "objectId": "shape1",
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Hello"}}]
                            }
                        },
                    }
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_get_presentation("pres1"))

    assert result["status"] == "success"
    assert result["title"] == "My Deck"
    assert result["slide_count"] == 1
    assert result["slides"][0]["text"] == ["Hello"]


def test_get_presentation_resolves_full_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "abc123",
        "title": "My Deck",
        "slides": [],
    }
    _mock_slides_service(monkeypatch, presentations)

    url = "https://docs.google.com/presentation/d/abc123/edit#slide=id.p"
    result = json.loads(google_slides.google_slides_get_presentation(url))

    assert result["status"] == "success"
    assert presentations.get.call_args.kwargs["presentationId"] == "abc123"


def test_get_presentation_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_get_presentation("pres1"))

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_create_presentation_returns_link_and_id(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "New Deck",
        "slides": [{"objectId": "p", "pageElements": []}],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("New Deck"))

    assert result["status"] == "success"
    assert result["presentation_id"] == "pres1"
    assert result["link"] == "https://docs.google.com/presentation/d/pres1/edit"
    assert result["slide_count"] == 1
    assert result["default_slide_id"] == "p"
    presentations.batchUpdate.assert_not_called()


def test_create_presentation_reads_initial_slide_when_create_omits_slides(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "New Deck",
    }
    _mock_slides_service(monkeypatch, presentations)
    presentations.get.return_value.execute.return_value = {
        "slides": [{"objectId": "p", "pageElements": []}]
    }

    result = json.loads(google_slides.google_slides_create_presentation("New Deck"))

    assert result["status"] == "success"
    assert result["slide_count"] == 1
    assert result["default_slide_id"] == "p"
    assert presentations.get.call_count == 1


def test_create_presentation_does_not_delete_nonempty_initial_slide(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "New Deck",
        "slides": [
            {
                "objectId": "p",
                "pageElements": [
                    _placeholder_element("p_title", "CENTERED_TITLE", "Existing")
                ],
            }
        ],
    }
    presentations.get.return_value.execute.return_value = {
        "slides": presentations.create.return_value.execute.return_value["slides"]
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("New Deck"))

    assert result["status"] == "success"
    assert result["slide_count"] == 1
    presentations.batchUpdate.assert_not_called()


def test_add_slide_removes_default_blank_slide_in_same_batch(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [{"objectId": "p", "pageElements": []}]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Title", body="Detail", default_slide_id="p"
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is True
    requests = _batch_update_requests(presentations)
    assert "createSlide" in requests[0]
    assert requests[-1] == {"deleteObject": {"objectId": "p"}}


def test_add_slide_removes_default_slide_across_mcp_processes(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "New Deck",
        "slides": [{"objectId": "p", "pageElements": []}],
    }
    presentations.get.return_value.execute.return_value = {
        "slides": [{"objectId": "p", "pageElements": []}]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    created = json.loads(google_slides.google_slides_create_presentation("New Deck"))
    google_slides._CREATED_DEFAULT_SLIDES.clear()
    added = json.loads(
        google_slides.google_slides_add_slide(
            created["presentation_id"], title="Title", body="Detail"
        )
    )

    assert added["default_slide_removed"] is True
    assert _batch_update_requests(presentations)[-1] == {
        "deleteObject": {"objectId": "p"}
    }


def test_add_slide_removes_known_default_slide_when_other_slides_exist(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {"objectId": "p", "pageElements": []},
            {
                "objectId": "existing-content",
                "pageElements": [
                    _placeholder_element("existing-title", "TITLE", "Existing")
                ],
            },
        ]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Title", body="Detail", default_slide_id="p"
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is True
    assert _batch_update_requests(presentations)[-1] == {
        "deleteObject": {"objectId": "p"}
    }


def test_add_slide_preserves_first_blank_slide_without_id_when_other_pages_exist(
    monkeypatch,
):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {"objectId": "p", "pageElements": []},
            {
                "objectId": "existing-content",
                "pageElements": [
                    _placeholder_element("existing-title", "TITLE", "Existing")
                ],
            },
        ]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Title", body="Detail")
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


def test_add_slide_keeps_untracked_intentional_blank_slide(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [{"objectId": "intentional-blank", "pageElements": []}]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Title",
            body="Detail",
            preserve_blank_slide=True,
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    requests = _batch_update_requests(presentations)
    assert not any("deleteObject" in request for request in requests)


def test_add_slide_keeps_image_only_existing_slide(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {
                "objectId": "image-slide",
                "pageElements": [
                    {"image": {"contentUrl": "https://example.com/a.png"}}
                ],
            }
        ]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Title", body="Detail")
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    requests = _batch_update_requests(presentations)
    assert not any("deleteObject" in request for request in requests)


def test_add_slide_keeps_background_only_existing_slide_without_id(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {
                "objectId": "background-slide",
                "pageElements": [],
                "pageProperties": {
                    "pageBackgroundFill": {
                        "solidFill": {"color": {"rgbColor": {"red": 1}}}
                    }
                },
            }
        ]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Title", body="Detail")
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


def test_add_slide_keeps_unfilled_placeholder_cover_without_id(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {
                "objectId": "cover-slide",
                "pageElements": [_placeholder_element("title", "TITLE")],
            }
        ]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Title", body="Detail")
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


def test_add_slide_preserve_blank_slide_does_not_delete_it_on_next_call(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "slides": [{"objectId": "default", "pageElements": []}]
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    google_slides._CREATED_DEFAULT_SLIDES["pres1"] = "default"

    preserved = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Title",
            body="Detail",
            default_slide_id="default",
            preserve_blank_slide=True,
        )
    )

    assert preserved["status"] == "success"
    assert "pres1" not in google_slides._CREATED_DEFAULT_SLIDES

    appended = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Second", body="More detail"
        )
    )

    assert appended["status"] == "success"
    assert appended["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


def _default_page_with_empty_placeholders():
    """A new presentation's first page as the Slides editor shows it: a
    CENTERED_TITLE and a SUBTITLE placeholder, both still empty."""
    return {
        "objectId": "p",
        "pageElements": [
            _placeholder_element("p_title", "CENTERED_TITLE"),
            _placeholder_element("p_subtitle", "SUBTITLE"),
        ],
    }


def _titled_default_page(title):
    """The default page after google_slides_create_presentation wrote the deck
    title into its CENTERED_TITLE placeholder."""
    return {
        "objectId": "p",
        "pageElements": [
            _placeholder_element("p_title", "CENTERED_TITLE", f"{title}\n"),
            _placeholder_element("p_subtitle", "SUBTITLE"),
        ],
    }


def test_outline_deck_maps_cover_and_body_placeholders_and_replaces_default_page(
    monkeypatch,
):
    """Create a deck from an outline the way the tool descriptions steer a
    caller: create, then a TITLE cover that passes default_slide_id, then a
    TITLE_AND_BODY slide with nested bullets. Each call is treated as its own
    MCP process, so nothing is remembered between them."""
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Q3 Review",
        "slides": [_default_page_with_empty_placeholders()],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    created = json.loads(google_slides.google_slides_create_presentation("Q3 Review"))

    assert created["default_slide_id"] == "p"
    assert _batch_update_requests(presentations) == [
        {"insertText": {"objectId": "p_title", "text": "Q3 Review"}}
    ]
    next_step = created["next_step"]
    assert next_step.startswith("When you add slides to this presentation")
    assert "default_slide_id='p'" in next_step
    assert "already shows the deck title" in next_step
    assert "Unless the user asked for a different first slide" in next_step
    assert "layout='TITLE'" in next_step
    assert "do not add a second cover" in next_step

    google_slides._CREATED_DEFAULT_SLIDES.clear()
    presentations.get.return_value.execute.return_value = {
        "title": "Q3 Review",
        "slides": [_titled_default_page("Q3 Review")],
    }
    cover = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Q3 Review",
            body="Prepared for the leadership team",
            layout="TITLE",
            default_slide_id=created["default_slide_id"],
        )
    )

    assert cover["status"] == "success"
    assert cover["default_slide_removed"] is True
    cover_requests = _batch_update_requests(presentations)
    create_cover = cover_requests[0]["createSlide"]
    assert create_cover["slideLayoutReference"] == {"predefinedLayout": "TITLE"}
    cover_title_id = _placeholder_object_id(create_cover, "CENTERED_TITLE")
    cover_subtitle_id = _placeholder_object_id(create_cover, "SUBTITLE")
    assert cover_requests[1:] == [
        {"insertText": {"objectId": cover_title_id, "text": "Q3 Review"}},
        {
            "insertText": {
                "objectId": cover_subtitle_id,
                "text": "Prepared for the leadership team",
            }
        },
        {"deleteObject": {"objectId": "p"}},
    ]

    google_slides._CREATED_DEFAULT_SLIDES.clear()
    presentations.get.return_value.execute.return_value = {
        "slides": [
            {
                "objectId": cover["slide_id"],
                "pageElements": [
                    _placeholder_element(cover_title_id, "CENTERED_TITLE", "Q3 Review"),
                    _placeholder_element(
                        cover_subtitle_id,
                        "SUBTITLE",
                        "Prepared for the leadership team",
                    ),
                ],
            }
        ]
    }
    body_slide = json.loads(
        google_slides.google_slides_add_slide(
            "pres1",
            title="Highlights",
            body="- Revenue up 12%\n  - Driven by renewals\n- Two new regions",
        )
    )

    assert body_slide["status"] == "success"
    assert body_slide["default_slide_removed"] is False
    body_requests = _batch_update_requests(presentations)
    create_body = body_requests[0]["createSlide"]
    assert create_body["slideLayoutReference"] == {"predefinedLayout": "TITLE_AND_BODY"}
    title_id = _placeholder_object_id(create_body, "TITLE")
    body_id = _placeholder_object_id(create_body, "BODY")
    assert body_requests[1:] == [
        {"insertText": {"objectId": title_id, "text": "Highlights"}},
        {
            "insertText": {
                "objectId": body_id,
                "text": "Revenue up 12%\n\tDriven by renewals\nTwo new regions",
            }
        },
        {
            "createParagraphBullets": {
                "objectId": body_id,
                "textRange": {"type": "ALL"},
                "bulletPreset": "BULLET_DISC_CIRCLE_SQUARE",
            }
        },
    ]


def test_create_presentation_writes_the_deck_title_into_the_default_page(
    monkeypatch,
):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Q3\nReview ",
        "slides": [_default_page_with_empty_placeholders()],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("Q3 Review"))

    assert result["status"] == "success"
    assert result["default_slide_id"] == "p"
    presentations.batchUpdate.assert_called_once()
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "pres1"
    assert _batch_update_requests(presentations) == [
        {"insertText": {"objectId": "p_title", "text": "Q3 Review"}}
    ]
    assert google_slides._CREATED_DEFAULT_SLIDES == {"pres1": "p"}


def test_create_presentation_still_succeeds_when_the_title_write_fails(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Q3 Review",
        "slides": [_default_page_with_empty_placeholders()],
    }
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("Q3 Review"))

    assert result["status"] == "success"
    assert result["default_slide_id"] == "p"
    assert "already shows the deck title" not in result["next_step"]
    assert "can remain as an empty first slide" in result["next_step"]


def test_add_slide_replaces_the_titled_default_page_for_a_different_first_slide(
    monkeypatch,
):
    """The user asked for an agenda as the first slide: the default page that
    create titled is still replaced in the same batch, so no stray title page
    is left in front of it."""
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "title": "Q3 Review",
        "slides": [_titled_default_page("Q3 Review")],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    google_slides._CREATED_DEFAULT_SLIDES["pres1"] = "p"

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Agenda", body="Results\nPlans", default_slide_id="p"
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is True
    requests = _batch_update_requests(presentations)
    assert "createSlide" in requests[0]
    assert requests[-1] == {"deleteObject": {"objectId": "p"}}
    assert "pres1" not in google_slides._CREATED_DEFAULT_SLIDES


@pytest.mark.parametrize(
    "remembered", [False, True], ids=["new-process", "same-process"]
)
def test_add_slide_without_id_keeps_the_titled_default_page_as_the_cover(
    monkeypatch, remembered
):
    """Without default_slide_id the titled default page stays as the cover,
    as next_step and the descriptions say, also when this process remembers
    the page from create."""
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "title": "Q3 Review",
        "slides": [_titled_default_page("Q3 Review")],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)
    if remembered:
        google_slides._CREATED_DEFAULT_SLIDES["pres1"] = "p"

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="Highlights", body="Up")
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )
    assert "pres1" not in google_slides._CREATED_DEFAULT_SLIDES


def test_add_slide_keeps_the_titled_default_page_once_other_pages_exist(
    monkeypatch,
):
    """A first call without the id kept the titled default page as the cover;
    passing the id on a later call must not delete the deck's only cover."""
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "title": "Q3 Review",
        "slides": [
            _titled_default_page("Q3 Review"),
            {
                "objectId": "s1",
                "pageElements": [
                    _placeholder_element("s1_title", "TITLE", "Highlights"),
                    _placeholder_element("s1_body", "BODY", "Up"),
                ],
            },
        ],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Plans", body="Next", default_slide_id="p"
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


@pytest.mark.parametrize(
    "elements",
    [
        pytest.param(
            [
                _placeholder_element("p_title", "CENTERED_TITLE", "Another title"),
                _placeholder_element("p_subtitle", "SUBTITLE"),
            ],
            id="title-changed",
        ),
        pytest.param(
            [
                _placeholder_element("p_title", "CENTERED_TITLE", "Q3 Review"),
                _placeholder_element("p_subtitle", "SUBTITLE", "Added by the user"),
            ],
            id="subtitle-added",
        ),
        pytest.param(
            [
                _placeholder_element("p_title", "CENTERED_TITLE", "Q3 Review"),
                {"objectId": "img", "image": {}},
            ],
            id="image-added",
        ),
    ],
)
def test_add_slide_keeps_a_default_page_that_was_edited_after_create(
    monkeypatch, elements
):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "title": "Q3 Review",
        "slides": [{"objectId": "p", "pageElements": elements}],
    }
    presentations.batchUpdate.return_value.execute.return_value = {}
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "pres1", title="Highlights", body="Up", default_slide_id="p"
        )
    )

    assert result["status"] == "success"
    assert result["default_slide_removed"] is False
    assert not any(
        "deleteObject" in request for request in _batch_update_requests(presentations)
    )


async def test_first_slide_steering_defers_to_the_users_own_first_slide():
    tools = {tool.name: tool for tool in await google_slides.mcp.list_tools()}

    for name, phrase in (
        ("google_slides_create_presentation", "Unless the user asked for"),
        ("google_slides_add_slide", "Unless the user specified"),
    ):
        description = " ".join(tools[name].description.split())
        assert "default_slide_id" in description
        assert f"{phrase} a different first slide" in description
        assert 'layout="TITLE"' in description


def test_create_presentation_omits_next_step_without_a_default_page(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "New Deck",
        "slides": [
            {
                "objectId": "p",
                "pageElements": [
                    _placeholder_element("p_title", "CENTERED_TITLE", "Existing")
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("New Deck"))

    assert result["default_slide_id"] is None
    assert "next_step" not in result


def test_import_pptx_converts_and_verifies_native_google_slides(monkeypatch, tmp_path):
    pptx_path = tmp_path / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "Title")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
        "mimeType": "application/vnd.google-apps.presentation",
        "webViewLink": "https://docs.google.com/presentation/d/pres1/edit",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Title\n"}}]
                            }
                        }
                    }
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_import_pptx(str(pptx_path), title="Designed Deck")
    )

    assert result["status"] == "success"
    assert result["presentation_id"] == "pres1"
    assert result["slide_count"] == 1
    request = drive.files.return_value.create.call_args.kwargs
    assert request["body"]["mimeType"] == "application/vnd.google-apps.presentation"
    assert request["media_body"]._mimetype == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )


def test_extract_pptx_slide_text_reads_title_and_body(tmp_path):
    from pptx import Presentation

    pptx_path = tmp_path / "designed-deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Detail one\nDetail two"
    presentation.save(pptx_path)

    assert google_slides._extract_pptx_slide_text(pptx_path) == [
        "Title Detail one Detail two"
    ]


def test_import_pptx_resolves_relative_generated_file_from_allowed_root(
    monkeypatch, tmp_path
):
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    pptx_path = output_dir / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "Title")
    unrelated_dir = tmp_path / "unrelated"
    unrelated_dir.mkdir()
    monkeypatch.chdir(unrelated_dir)
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Title\n"}}]
                            }
                        }
                    }
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_import_pptx("output/designed-deck.pptx")
    )

    assert result["status"] == "success"
    assert drive.files.return_value.create.call_args.kwargs[
        "media_body"
    ]._fd.name == str(pptx_path)


def test_import_pptx_rejects_missing_text_after_conversion(monkeypatch, tmp_path):
    pptx_path = tmp_path / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "Title Detail")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Title\n"}}]
                            }
                        }
                    }
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "validation_failed"
    assert result["empty_slide_numbers"] == []
    assert result["missing_text_slide_numbers"] == [1]


def test_import_pptx_allows_drive_to_reorder_text_elements(monkeypatch, tmp_path):
    pptx_path = tmp_path / "reordered-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "Title Detail")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Reordered Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Reordered Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Detail\n"}}]
                            }
                        }
                    },
                    {
                        "shape": {
                            "text": {
                                "textElements": [{"textRun": {"content": "Title\n"}}]
                            }
                        }
                    },
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "success"


def test_import_pptx_rejects_empty_slide_after_conversion(monkeypatch, tmp_path):
    pptx_path = tmp_path / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "Title")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [{"objectId": "slide1", "pageElements": []}],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "validation_failed"
    assert result["empty_slide_numbers"] == [1]
    assert "presentation_id" not in result
    assert "link" not in result
    drive.files.return_value.delete.assert_called_once_with(
        fileId="pres1", supportsAllDrives=True
    )


def test_element_text_reads_grouped_shapes():
    element = {
        "elementGroup": {
            "children": [
                {
                    "shape": {
                        "text": {
                            "textElements": [{"textRun": {"content": "Grouped text"}}]
                        }
                    }
                }
            ]
        }
    }

    assert google_slides._element_text(element) == "Grouped text"


def test_element_text_reads_auto_text_content():
    element = {
        "shape": {
            "text": {
                "textElements": [{"autoText": {"type": "SLIDE_NUMBER", "content": "3"}}]
            }
        }
    }

    assert google_slides._element_text(element) == "3"


def test_import_pptx_allows_intentional_blank_source_slide(monkeypatch, tmp_path):
    pptx_path = tmp_path / "blank-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Blank Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Blank Deck",
        "slides": [{"objectId": "slide1", "pageElements": []}],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "success"


def test_import_pptx_allows_image_only_slide_after_conversion(monkeypatch, tmp_path):
    pptx_path = tmp_path / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "")
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [
            {
                "objectId": "slide1",
                "pageElements": [
                    {"image": {"contentUrl": "https://example.com/a.png"}}
                ],
            }
        ],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "success"
    assert result["empty_slide_numbers"] == []


def test_import_pptx_rejects_dropped_visual_only_slide(monkeypatch, tmp_path):
    pptx_path = tmp_path / "visual-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch, "", has_content=[True])
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Visual Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Visual Deck",
        "slides": [{"objectId": "slide1", "pageElements": []}],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "validation_failed"
    assert result["empty_slide_numbers"] == [1]
    drive.files.return_value.delete.assert_called_once_with(
        fileId="pres1", supportsAllDrives=True
    )


def test_import_pptx_rejects_conversion_with_no_slides(monkeypatch, tmp_path):
    pptx_path = tmp_path / "designed-deck.pptx"
    pptx_path.write_bytes(b"pptx-bytes")
    _mock_pptx_text(monkeypatch)
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    drive = Mock()
    drive.files.return_value.create.return_value.execute.return_value = {
        "id": "pres1",
        "name": "Designed Deck",
    }
    monkeypatch.setattr(google_slides, "get_drive_service", lambda: drive)

    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "pres1",
        "title": "Designed Deck",
        "slides": [],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_import_pptx(str(pptx_path)))

    assert result["status"] == "validation_failed"
    assert result["slide_count"] == 0
    assert result["empty_slide_numbers"] == []


def test_import_pptx_rejects_path_outside_allowlist(monkeypatch, tmp_path):
    monkeypatch.setenv("XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS", str(tmp_path))

    result = json.loads(
        google_slides.google_slides_import_pptx("/private/tmp/not-allowed.pptx")
    )

    assert result["status"] == "error"
    assert "outside the allowed directories" in result["message"]


def test_resolve_pptx_upload_path_rejects_oversized_file(monkeypatch, tmp_path):
    monkeypatch.setattr(google_slides, "allowed_dirs_from_env", lambda _: [tmp_path])
    oversized = tmp_path / "oversized.pptx"
    with oversized.open("wb") as file_handle:
        file_handle.truncate(google_slides._MAX_PPTX_UPLOAD_BYTES + 1)

    with pytest.raises(ValueError, match="100 MB"):
        google_slides._resolve_pptx_upload_path(str(oversized))


def test_resolve_pptx_upload_path_surfaces_symlink_loop(monkeypatch, tmp_path):
    monkeypatch.setattr(google_slides, "allowed_dirs_from_env", lambda _: [tmp_path])
    loop_path = tmp_path / "loop.pptx"
    loop_path.symlink_to(loop_path)

    def _raise_symlink_loop(_path):
        raise RuntimeError("symlink loop")

    monkeypatch.setattr(Path, "resolve", _raise_symlink_loop)

    with pytest.raises(OSError) as exc_info:
        google_slides._resolve_pptx_upload_path(str(loop_path))
    assert exc_info.value.errno == errno.ELOOP


def test_create_presentation_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.create.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_create_presentation("New Deck"))

    assert result["status"] == "error"
    assert "boom" in result["message"]


def test_batch_update_forwards_requests_and_returns_replies(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {
        "presentationId": "pres1",
        "replies": [{"createShape": {"objectId": "shape1"}}],
    }
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_batch_update(
            "pres1", '[{"createShape": {"objectId": "shape1"}}]'
        )
    )

    assert result["status"] == "success"
    assert result["replies"] == [{"createShape": {"objectId": "shape1"}}]
    sent_requests = presentations.batchUpdate.call_args.kwargs["body"]["requests"]
    assert sent_requests == [{"createShape": {"objectId": "shape1"}}]


def test_batch_update_rejects_non_list_json(monkeypatch):
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_batch_update("pres1", '{"not": "a list"}')
    )

    assert result["status"] == "error"
    presentations.batchUpdate.assert_not_called()


def test_batch_update_rejects_malformed_json(monkeypatch):
    """Not merely wrong-shaped JSON (a dict instead of a list, covered
    above) but genuinely invalid JSON syntax — must surface as a clean
    error payload rather than an unhandled JSONDecodeError."""
    presentations = Mock()
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_batch_update("pres1", "{not valid json")
    )

    assert result["status"] == "error"
    presentations.batchUpdate.assert_not_called()


def test_batch_update_resolves_full_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.return_value = {
        "presentationId": "abc123",
        "replies": [],
    }
    _mock_slides_service(monkeypatch, presentations)

    url = "https://docs.google.com/presentation/d/abc123/edit#slide=id.p"
    result = json.loads(google_slides.google_slides_batch_update(url, "[]"))

    assert result["status"] == "success"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "abc123"


def test_batch_update_returns_error_payload_on_api_failure(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = RuntimeError("boom")
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_batch_update("pres1", "[]"))

    assert result["status"] == "error"
    assert "boom" in result["message"]


class _HttpResponse:
    def __init__(self, status: int):
        self.status = status
        self.reason = "error"


def _http_error(status: int, body: dict):
    from googleapiclient.errors import HttpError

    return HttpError(
        _HttpResponse(status),
        json.dumps(body).encode("utf-8"),
        uri="https://slides.googleapis.com/v1/presentations/pres1?alt=json",
    )


def test_get_presentation_maps_not_found_to_an_actionable_message(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.side_effect = _http_error(
        404,
        {"error": {"code": 404, "message": "Requested entity was not found."}},
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_get_presentation("pres1"))

    assert result["status"] == "error"
    message = result["message"]
    assert message.startswith("Google Slides could not open this presentation")
    assert "google_slides_create_presentation" in message
    assert message.endswith("HTTP 404 Requested entity was not found.")


_VIEW_ONLY_403 = {
    "error": {
        "code": 403,
        "message": "The caller does not have permission",
        "status": "PERMISSION_DENIED",
    }
}


def _assert_refused_edit(result):
    assert result["status"] == "error"
    message = result["message"]
    assert message.startswith("Google Slides could not make this change")
    assert "only view or comment access" in message
    assert "can edit the presentation" in message
    assert "could not open" not in message
    assert message.endswith("HTTP 403 The caller does not have permission")


def test_add_slide_explains_a_permission_denied_read(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.side_effect = _http_error(
        403, _VIEW_ONLY_403
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide("pres1", title="T", body="B")
    )

    _assert_refused_edit(result)
    assert "may not be able to open it at all" in result["message"]
    presentations.batchUpdate.assert_not_called()


@pytest.mark.parametrize(
    "invoke",
    [
        pytest.param(
            lambda: google_slides.google_slides_add_slide("pres1", title="T", body="B"),
            id="add_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_update_slide(
                "pres1", "slide1", title="T"
            ),
            id="update_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_delete_slide("pres1", "slide1"),
            id="delete_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_batch_update(
                "pres1", '[{"deleteObject": {"objectId": "slide1"}}]'
            ),
            id="batch_update",
        ),
    ],
)
def test_editing_tools_explain_a_refused_edit_on_a_readable_presentation(
    monkeypatch, invoke
):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = _http_error(
        403, _VIEW_ONLY_403
    )
    _mock_slides_service(monkeypatch, presentations)
    _mock_presentation_get(
        presentations, "slide1", [_placeholder_element("title_obj", "TITLE")]
    )

    result = json.loads(invoke())

    _assert_refused_edit(result)
    presentations.batchUpdate.assert_called_once()


def test_editing_tool_maps_not_found_to_an_actionable_message(monkeypatch):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = _http_error(
        404,
        {"error": {"code": 404, "message": "Requested entity was not found."}},
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_batch_update("pres1", "[]"))

    assert result["message"].startswith(
        "Google Slides could not open this presentation"
    )


async def test_get_presentation_description_says_drive_is_not_needed():
    tools = {tool.name: tool for tool in await google_slides.mcp.list_tools()}

    description = " ".join(tools["google_slides_get_presentation"].description.split())
    assert "cannot search for or list presentations by name" in description
    assert "Connecting Google Drive is not needed" in description
    assert (
        "use google_drive_search to find its id if that tool is available"
        in description
    )
    assert "otherwise, or if it finds nothing, ask the user to paste the link" in (
        description
    )


def test_get_presentation_rejects_a_title_without_calling_the_api(monkeypatch):
    get_service = Mock()
    monkeypatch.setattr(google_slides, "get_slides_service", get_service)

    result = json.loads(google_slides.google_slides_get_presentation("Sales kickoff"))

    assert result["status"] == "error"
    message = result["message"]
    assert "'Sales kickoff' is not a Google Slides link" in message
    assert "cannot search for or list presentations by name" in message
    assert "https://docs.google.com/presentation/d/" in message
    get_service.assert_not_called()


def test_get_presentation_rejects_a_published_link_instead_of_reading_e_as_the_id(
    monkeypatch,
):
    get_service = Mock()
    monkeypatch.setattr(google_slides, "get_slides_service", get_service)

    result = json.loads(
        google_slides.google_slides_get_presentation(
            "https://docs.google.com/presentation/d/e/2PACX-1vRabc123/pub?start=false"
        )
    )

    assert result["status"] == "error"
    assert "is a published-to-the-web link" in result["message"]
    get_service.assert_not_called()


def test_get_presentation_resolves_multi_account_presentation_url(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.return_value = {
        "presentationId": "abc123",
        "title": "My Deck",
        "slides": [],
    }
    _mock_slides_service(monkeypatch, presentations)

    url = "https://docs.google.com/presentation/u/1/d/abc123/edit"
    result = json.loads(google_slides.google_slides_get_presentation(url))

    assert result["status"] == "success"
    assert presentations.get.call_args.kwargs["presentationId"] == "abc123"


# The Sheets API's Office-file 400 (its first sentence). The Slides tools
# use the same check, in case the Slides API answers a PowerPoint file
# stored in Drive the same way.
_OFFICE_FILE_400 = {
    "error": {
        "code": 400,
        "message": "This operation is not supported for this document",
        "status": "FAILED_PRECONDITION",
    }
}


def test_get_presentation_explains_a_powerpoint_file_stored_in_drive(monkeypatch):
    presentations = Mock()
    presentations.get.return_value.execute.side_effect = _http_error(
        400, _OFFICE_FILE_400
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(google_slides.google_slides_get_presentation("pres1"))

    assert result["status"] == "error"
    message = result["message"]
    # Google Slides itself opens the file in Office compatibility mode; only
    # its API cannot.
    assert message.startswith(
        "The Google Slides tools cannot open this file: it is most likely a "
        "PowerPoint file (.pptx) stored in Google Drive rather than a Google "
        "Slides presentation. Google Slides can open such a file in Office "
        "compatibility mode"
    )
    assert "read the downloaded copy with read_pptx" in message
    assert "File > Save as Google Slides" in message
    assert "is a Google Docs or Google Sheets file instead" in message
    assert message.endswith(
        "Google API response: HTTP 400 This operation is not supported for this "
        "document"
    )


# Every Slides tool that changes an existing presentation, with the request
# that gets the 400: each must pass editing=True to the error message too,
# not only to the id resolver.
@pytest.mark.parametrize(
    ("call", "request_of"),
    [
        pytest.param(
            lambda: google_slides.google_slides_add_slide("pres1", title="T", body="B"),
            lambda presentations: presentations.get,
            id="add_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_update_slide(
                "pres1", "slide1", title="T"
            ),
            lambda presentations: presentations.get,
            id="update_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_delete_slide("pres1", "slide1"),
            lambda presentations: presentations.get,
            id="delete_slide",
        ),
        pytest.param(
            lambda: google_slides.google_slides_batch_update(
                "pres1", '[{"deleteObject": {"objectId": "shape1"}}]'
            ),
            lambda presentations: presentations.batchUpdate,
            id="batch_update",
        ),
    ],
)
def test_editing_tools_explain_a_powerpoint_file_stored_in_drive(
    monkeypatch, call, request_of
):
    presentations = Mock()
    request = request_of(presentations)
    request.return_value.execute.side_effect = _http_error(400, _OFFICE_FILE_400)
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(call())

    assert result["status"] == "error"
    assert request.call_args.kwargs["presentationId"] == "pres1"
    message = result["message"]
    assert message.startswith(
        "The Google Slides tools cannot open this file: it is most likely a "
        "PowerPoint file"
    )
    assert "To make this change with these tools" in message
    assert (
        "this creates a new presentation: the change will be made in that copy"
        in message
    )
    assert "read_pptx" not in message
    assert message.endswith(
        "Google API response: HTTP 400 This operation is not supported for this "
        "document"
    )


@pytest.mark.parametrize(
    ("call", "error_body"),
    [
        # The Slides API answers the id of an Excel file with a 400 that says
        # nothing about the file.
        pytest.param(
            lambda: google_slides.google_slides_get_presentation("pres1"),
            {
                "error": {
                    "code": 400,
                    "message": "Request contains an invalid argument.",
                    "status": "INVALID_ARGUMENT",
                }
            },
            id="get_presentation-invalid-argument",
        ),
        # A request built by the caller can fail a precondition for reasons
        # that have nothing to do with an Office file.
        pytest.param(
            lambda: google_slides.google_slides_batch_update(
                "pres1", '[{"deleteObject": {"objectId": "shape1"}}]'
            ),
            {
                "error": {
                    "code": 400,
                    "message": "Precondition check failed.",
                    "status": "FAILED_PRECONDITION",
                }
            },
            id="batch_update-failed-precondition",
        ),
    ],
)
def test_tools_keep_the_raw_error_for_another_400(monkeypatch, call, error_body):
    presentations = Mock()
    error = _http_error(400, error_body)
    presentations.get.return_value.execute.side_effect = error
    presentations.batchUpdate.return_value.execute.side_effect = error
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(call())

    assert result["status"] == "error"
    assert result["message"] == str(error)


@pytest.mark.parametrize(
    ("call", "editing"),
    [
        pytest.param(
            google_slides.google_slides_get_presentation, False, id="get_presentation"
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_add_slide(
                pres, title="T", body="B"
            ),
            True,
            id="add_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_update_slide(
                pres, "slide1", title="T"
            ),
            True,
            id="update_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_delete_slide(pres, "slide1"),
            True,
            id="delete_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_batch_update(pres, "[]"),
            True,
            id="batch_update",
        ),
    ],
)
def test_tools_give_their_own_next_steps_for_a_drive_file_link(
    monkeypatch, call, editing
):
    get_service = Mock()
    monkeypatch.setattr(google_slides, "get_slides_service", get_service)

    result = json.loads(call("https://drive.google.com/file/d/pres1/view"))

    assert result["status"] == "error"
    message = result["message"]
    assert "is a Google Drive file link, not a Google Slides link" in message
    assert "File > Save as Google Slides" in message
    # A tool that changes the file is not sent to read a downloaded copy, and
    # is told that saving as a Google Slides file makes a new copy.
    assert ("read the downloaded copy with read_pptx" in message) is not editing
    assert (
        "this creates a new presentation: the change will be made in that copy"
        in message
    ) is editing
    get_service.assert_not_called()


# A 400 whose message names a cause is about the request, so it keeps the
# raw error for an id taken from a Drive open?id= or uc?id= link too.
@pytest.mark.parametrize(
    "link",
    [
        "https://drive.google.com/open?id=pres1",
        "https://drive.google.com/uc?export=download&id=pres1",
    ],
)
@pytest.mark.parametrize(
    "error_body",
    [
        pytest.param(
            {
                "error": {
                    "code": 400,
                    "message": "Precondition check failed.",
                    "status": "FAILED_PRECONDITION",
                }
            },
            id="failed-precondition",
        ),
        pytest.param(
            {
                "error": {
                    "code": 400,
                    "message": (
                        "Invalid requests[0].deleteObject: The object (shape1) "
                        "could not be found."
                    ),
                    "status": "INVALID_ARGUMENT",
                }
            },
            id="request-field",
        ),
    ],
)
def test_batch_update_keeps_the_raw_error_for_a_400_with_a_cause_for_a_drive_open_or_uc_link(
    monkeypatch, link, error_body
):
    presentations = Mock()
    error = _http_error(400, error_body)
    presentations.batchUpdate.return_value.execute.side_effect = error
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_batch_update(
            link, '[{"deleteObject": {"objectId": "shape1"}}]'
        )
    )

    assert result["status"] == "error"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "pres1"
    assert result["message"] == str(error)


# A Drive open?id= or uc?id= link can name a native presentation or an
# uploaded file, so its id goes to the API. The Slides API answers the id of
# an uploaded file (an Excel file, at least) with a 400 that does not name
# the cause, so every tool must pass its own argument to the error message,
# which then adds the next steps for an uploaded file.
@pytest.mark.parametrize(
    "link",
    [
        "https://drive.google.com/open?id=pres1",
        "https://drive.google.com/uc?export=download&id=pres1",
    ],
)
@pytest.mark.parametrize(
    ("call", "request_of", "editing"),
    [
        pytest.param(
            google_slides.google_slides_get_presentation,
            lambda presentations: presentations.get,
            False,
            id="get_presentation",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_add_slide(
                pres, title="T", body="B"
            ),
            lambda presentations: presentations.get,
            True,
            id="add_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_update_slide(
                pres, "slide1", title="T"
            ),
            lambda presentations: presentations.get,
            True,
            id="update_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_delete_slide(pres, "slide1"),
            lambda presentations: presentations.get,
            True,
            id="delete_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_batch_update(
                pres, '[{"deleteObject": {"objectId": "shape1"}}]'
            ),
            lambda presentations: presentations.batchUpdate,
            True,
            id="batch_update",
        ),
    ],
)
def test_tools_add_next_steps_to_a_400_for_a_drive_open_or_uc_link(
    monkeypatch, link, call, request_of, editing
):
    presentations = Mock()
    request = request_of(presentations)
    request.return_value.execute.side_effect = _http_error(
        400,
        {
            "error": {
                "code": 400,
                "message": "Request contains an invalid argument.",
                "status": "INVALID_ARGUMENT",
            }
        },
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(call(link))

    assert result["status"] == "error"
    assert request.call_args.kwargs["presentationId"] == "pres1"
    message = result["message"]
    assert message.startswith(
        "Google Slides rejected this request (Google API response: HTTP 400 "
        "Request contains an invalid argument.). The presentation id was taken "
        "from a Google Drive link"
    )
    assert "File > Save as Google Slides" in message
    assert ("read the downloaded copy with read_pptx" in message) is not editing
    assert ("To make this change with these tools" in message) is editing
    assert message.endswith("the error is about the request itself.")


_NO_CAUSE_400 = {
    "error": {
        "code": 400,
        "message": "Request contains an invalid argument.",
        "status": "INVALID_ARGUMENT",
    }
}


# Once a tool has read the presentation, it is a Google Slides presentation
# whatever link named it, so a 400 for the tool's later request keeps
# Google's own error, even one that names no cause.
@pytest.mark.parametrize(
    "link",
    [
        "https://drive.google.com/open?id=pres1",
        "https://drive.google.com/uc?export=download&id=pres1",
    ],
)
@pytest.mark.parametrize(
    "call",
    [
        pytest.param(
            lambda pres: google_slides.google_slides_add_slide(
                pres, title="T", body="B"
            ),
            id="add_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_update_slide(
                pres, "slide1", title="T"
            ),
            id="update_slide",
        ),
        pytest.param(
            lambda pres: google_slides.google_slides_delete_slide(pres, "slide1"),
            id="delete_slide",
        ),
    ],
)
def test_tools_keep_the_raw_400_after_the_presentation_opened(monkeypatch, link, call):
    presentations = Mock()
    _mock_presentation_get(
        presentations,
        "slide1",
        [_placeholder_element("title_obj", "TITLE", text="Old title")],
    )
    error = _http_error(400, _NO_CAUSE_400)
    presentations.batchUpdate.return_value.execute.side_effect = error
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(call(link))

    assert result["status"] == "error"
    assert presentations.get.call_args.kwargs["presentationId"] == "pres1"
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "pres1"
    assert result["message"] == str(error)


# add_slide reads the presentation only when it may remove a default page.
# Without that read nothing shows which file the link names, so the 400
# still gets the next steps.
def test_add_slide_adds_next_steps_to_a_400_when_it_did_not_read_the_presentation(
    monkeypatch,
):
    presentations = Mock()
    presentations.batchUpdate.return_value.execute.side_effect = _http_error(
        400, _NO_CAUSE_400
    )
    _mock_slides_service(monkeypatch, presentations)

    result = json.loads(
        google_slides.google_slides_add_slide(
            "https://drive.google.com/open?id=pres1",
            title="T",
            body="B",
            preserve_blank_slide=True,
        )
    )

    assert result["status"] == "error"
    presentations.get.assert_not_called()
    assert presentations.batchUpdate.call_args.kwargs["presentationId"] == "pres1"
    assert result["message"].startswith(
        "Google Slides rejected this request (Google API response: HTTP 400 "
        "Request contains an invalid argument.). The presentation id was taken "
        "from a Google Drive link"
    )
    assert "To make this change with these tools" in result["message"]
